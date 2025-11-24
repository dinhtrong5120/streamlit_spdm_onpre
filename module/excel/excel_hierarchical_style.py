from openpyxl.styles import PatternFill, Font, Border, Side, Alignment, Color
from dataclasses import dataclass


# ----------------------------------- Meta -----------------------------------

# Meta Info
META_COL = 3
META_ROW = 6

# MetaColor
META_COLOR = PatternFill(fill_type='solid', fgColor='ffddebf7')

META_BORDER = Border(
    left=Side(style='medium'),
    right=Side(style='medium'),
    top=Side(style='medium'),
    bottom=Side(style='medium')
)


@dataclass
class MetaDef:
    col: int
    row: int
    value: str
    side: Side = META_BORDER
    fill: PatternFill = META_COLOR


PROJECT_COL = META_COL
PT_TYPE_COL = META_COL
LOT_COL = META_COL + 2
PHASE_COL = META_COL + 2

# meta title
PROJECT = 'Project'
PT_TYPE = 'PTタイプ'
LOT = 'Lot'
HIERARCHY = '階層'
WP = '性能'
PHASE = 'フェーズ'

META_DEF = {
    "project": MetaDef(col=PROJECT_COL,row=META_ROW,value=PROJECT),
    "pt_type": MetaDef(col=PT_TYPE_COL,row=META_ROW+1,value=PT_TYPE),
    "lot": MetaDef(col=LOT_COL,row=META_ROW,value=LOT),
    "phase": MetaDef(col=PHASE_COL,row=META_ROW+1,value=PHASE),
}


# ----------------------------------- Column Name -----------------------------------

# Sub Header Value
REQ_HEADER_VALUE = 'Requirement'
FUNC_HEADER_VALUE = 'Function'
LOGIC_HEADER_VALUE = 'Logic'
SENDER_HEADER_VALUE = '出し手'
RECEIVER_HEADER_VALUE = '受け手'

# Requirement
WP = '性能'
REQ_ITEM = '要求項目'
REQ_VALUE = '要求値'
REQ_UNIT = '単位'
REQ_CONDITION = '環境・運転条件'

# Function
FUNC_ITEM = '機能'
FUNC_VALUE = '機能目標'
FUNC_UNIT = '単位'

# Logic
LOGIC_ITEM = '要求'
LOGIC_VALUE = '要求値'
LOGIC_UNIT = '単位'
LOGIC_CONDITION = '環境・使用条件'
LOGIC_NOTE = 'Note'

# Sender
SENDER_APPROVAL = '承認'
SENDER_APPROVER = '承認者'
SENDER_DATE = '日付'
SENDER_COMMENT = 'コメント'

# Receiver
RECEIVER_APPROVAL = '承認'
RECEIVER_APPROVER = '承認者'
RECEIVER_DATE = '日付'
RECEIVER_COMMENT = 'コメント'

# -------------------------------------- Color --------------------------------------

VEHICLE_FILL = PatternFill(fill_type='solid', fgColor='FF002060')
SYSTEM_FILL = PatternFill(fill_type='solid', fgColor='FFED7D31')
UNIT_FILL = PatternFill(fill_type='solid', fgColor='FF00B0F0')
REQ_FILL = PatternFill(fill_type='solid', fgColor='FFDDEBF7')
FUNC_FILL = PatternFill(fill_type='solid', fgColor='FFFFFF99')
LOGIC_FILL = PatternFill(fill_type='solid', fgColor='FFFFCCFF')
SENDER_FILL = REQ_FILL
RECEIVER_FILL = REQ_FILL


# -------------------------------------- Font --------------------------------------

HR_FONT = Font(color='FFFFFFFF',name='Meiryo UI', size=18)
STANDARD_FONT = Font(color='00000000',name='Meiryo UI', size=9)


# -------------------------------------- Border --------------------------------------

# ボーダー定義
HEADER_BORDER = Border(
    left=Side(style='medium'),
    right=Side(style='medium'),
    top=Side(style='medium'),
    bottom=Side(style='thin')
)

COLUMN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='medium')
)

COLUMN1_BORDER = Border(
    left=Side(style='medium'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='medium')
)

LAST_COLUMN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='medium'),
    top=Side(style='thin'),
    bottom=Side(style='medium')
)


# ------------------------------------- Drop Column -------------------------------------

# DROPするカラム
DROP_COLUMN = [
    "{0}r_pj_id",
    "{0}rfl_id",
    "{0}phase_id",
    "{0}phase",
    "{0}r_wp",
    "{0}flag_to",
    "{0}to_solving_value",
    "{0}flag_display_on_summary_logic",
    "{0}r_wp_id",
    "{0}l_wp",
    "{0}r_s_id",
    "{0}f_id",
    "{0}r_item_2",
    "{0}l_s_id",
    "{0}l_s_idc_f_id",
    "{0}user_memo",
    "c_related_se_parameter_id",
    "{0}r_item_index",
    "{0}user_memoc_related_se_parameter_id",
    "{0}allocation",
    "{0}index",
    "rfl_id",
    "ph_id",
    "{0}sender_selected",
    "{0}receiver_selected",
    "wp",
    "project_code",
    "lot",
    "archi",
    "phase",
    "hierarchy",
    "{0}n_wp"
]


# ------------------------------------- Static Def -------------------------------------

# 初期列
INITIAL_COL = 3

# 階層表示からの初期行
INITIAL_ROW = 9

# column width limit
LIMIT_WIDTH = 40

# gird間のマージン
MARGIN_COL = 2

# Row Def
HIERARCHY_ROW = INITIAL_ROW
HEADER_ROW = INITIAL_ROW + 2
COLUMN_ROW = HEADER_ROW + 1
VALUE_ROW = COLUMN_ROW + 1

# 空白が入る場合、罫線を削除する対象の列
BORDER_DELETE_TARGET_COL = ['wp','r_item','req','r_scene','f_item','func','l_item','logic','l_scene']

# ------------------------------------- DataClass -------------------------------------

@dataclass
class HeaderDef:
    col: int
    value: str
    max_col: int
    fill: PatternFill
    side: Side = HEADER_BORDER
    row: int = HEADER_ROW

@dataclass
class ColumnDef:
    col: int
    value: str
    fill: PatternFill
    side: Side = COLUMN_BORDER
    row: int = COLUMN_ROW


# ------------------------------------- Grid Class -------------------------------------

class GridBlock:

    def __init__(self,init_col:int,grid_width=21):
        """グリッド定義作成

        Args:
            init_col (int): グリッド開始列
            grid_width (int, optional): グリッド長さ. Defaults to 21.
        """

        self.init_col = init_col
        self.grid_width = grid_width

        # 階層ヘッダーの範囲(=1グリッドの長さ) 0はじまりである為、-1
        HR_MAX_COL = self.init_col + grid_width - 1

        # ------------------------------------- Header -------------------------------------

        # Req~Receiver Header Location
        REQ_HEADER = self.init_col
        FUNC_HEADER = self.init_col + 5
        LOGIC_HEADER = self.init_col + 8
        SENDER_HEADER = self.init_col + 13
        RECEIVER_HEADER = self.init_col + 17

        # ------------------------------------- Column -------------------------------------

        # WP
        WP_COL= self.init_col

        # Requirement
        REQ_ITEM_COL = self.init_col + 1
        REQ_VALUE_COL = self.init_col + 2
        REQ_UNIT_COL = self.init_col + 3
        REQ_CONDITION_COL = self.init_col + 4

        # Function
        FUNC_ITEM_COL = self.init_col + 5
        FUNC_VALUE_COL = self.init_col + 6
        FUNC_UNIT_COL = self.init_col + 7

        # Logic
        LOGIC_ITEM_COL = self.init_col + 8
        LOGIC_VALUE_COL = self.init_col + 9
        LOGIC_UNIT_COL = self.init_col + 10
        LOGIC_CONDITION_COL = self.init_col + 11
        LOGIC_NOTE_COL = self.init_col + 12

        # Sender
        SENDER_APPROVAL_COL = self.init_col + 13
        SENDER_APPROVER_COL = self.init_col + 14
        SENDER_DATE_COL = self.init_col + 15
        SENDER_COMMENT_COL = self.init_col + 16

        # Receiver
        RECEIVER_APPROVAL_COL = self.init_col + 17
        RECEIVER_APPROVER_COL = self.init_col + 18
        RECEIVER_DATE_COL = self.init_col + 19
        RECEIVER_COMMENT_COL = self.init_col + 20


        # ------------------------------------- Column Def -------------------------------------

        self.column_defs = {
            "wp": ColumnDef(col=WP_COL,value=WP,fill=REQ_FILL,side=COLUMN1_BORDER),
            "req_col1": ColumnDef(col=REQ_ITEM_COL,value=REQ_ITEM,fill=REQ_FILL),
            "req_col2": ColumnDef(col=REQ_VALUE_COL,value=REQ_VALUE,fill=REQ_FILL),
            "req_col3": ColumnDef(col=REQ_UNIT_COL,value=REQ_UNIT,fill=REQ_FILL),
            "req_col4": ColumnDef(col=REQ_CONDITION_COL,value=REQ_CONDITION,fill=REQ_FILL,side=LAST_COLUMN_BORDER),
            "func_col1": ColumnDef(col=FUNC_ITEM_COL,value=FUNC_ITEM,fill=FUNC_FILL),
            "func_col2": ColumnDef(col=FUNC_VALUE_COL,value=FUNC_VALUE,fill=FUNC_FILL),
            "func_col3": ColumnDef(col=FUNC_UNIT_COL,value=FUNC_UNIT,fill=FUNC_FILL,side=LAST_COLUMN_BORDER),
            "logic_col1": ColumnDef(col=LOGIC_ITEM_COL,value=LOGIC_ITEM,fill=LOGIC_FILL),
            "logic_col2": ColumnDef(col=LOGIC_VALUE_COL,value=LOGIC_VALUE,fill=LOGIC_FILL),
            "logic_col3": ColumnDef(col=LOGIC_UNIT_COL,value=LOGIC_UNIT,fill=LOGIC_FILL),
            "logic_col4": ColumnDef(col=LOGIC_CONDITION_COL,value=LOGIC_CONDITION,fill=LOGIC_FILL),
            "logic_col5": ColumnDef(col=LOGIC_NOTE_COL,value=LOGIC_NOTE,fill=LOGIC_FILL,side=LAST_COLUMN_BORDER),
            "sender_col1": ColumnDef(col=SENDER_APPROVAL_COL,value=SENDER_APPROVAL,fill=REQ_FILL),
            "sender_col2": ColumnDef(col=SENDER_APPROVER_COL,value=SENDER_APPROVER,fill=REQ_FILL),
            "sender_col3": ColumnDef(col=SENDER_DATE_COL,value=SENDER_DATE,fill=REQ_FILL),
            "sender_col4": ColumnDef(col=SENDER_COMMENT_COL,value=SENDER_COMMENT,fill=REQ_FILL,side=LAST_COLUMN_BORDER),
            "receiver_col1": ColumnDef(col=RECEIVER_APPROVAL_COL,value=RECEIVER_APPROVAL,fill=REQ_FILL),
            "receiver_col2": ColumnDef(col=RECEIVER_APPROVER_COL,value=RECEIVER_APPROVER,fill=REQ_FILL),
            "receiver_col3": ColumnDef(col=RECEIVER_DATE_COL,value=RECEIVER_DATE,fill=REQ_FILL),
            "receiver_col4": ColumnDef(col=RECEIVER_COMMENT_COL,value=RECEIVER_COMMENT,fill=REQ_FILL,side=LAST_COLUMN_BORDER)
        }

        # 各カラムの区切りを定義
        self.medium_border_cols = [
            0,FUNC_HEADER - self.init_col,LOGIC_HEADER - self.init_col, SENDER_HEADER - self.init_col, RECEIVER_HEADER - 3
        ]

        # ------------------------------------- Header Def -------------------------------------

        self.header_defs = {
            "req_header": HeaderDef(col=self.init_col,value=REQ_HEADER_VALUE,max_col=FUNC_HEADER - 1,fill=REQ_FILL),
            "func_header": HeaderDef(col=FUNC_HEADER,value=FUNC_HEADER_VALUE,max_col=LOGIC_HEADER - 1,fill=FUNC_FILL),
            "logic_header": HeaderDef(col=LOGIC_HEADER,value=LOGIC_HEADER_VALUE,max_col=SENDER_HEADER - 1,fill=LOGIC_FILL),
            "sender_header": HeaderDef(col=SENDER_HEADER,value=SENDER_HEADER_VALUE,max_col=RECEIVER_HEADER - 1,fill=SENDER_FILL),
            "receiver_header": HeaderDef(col=RECEIVER_HEADER,value=RECEIVER_HEADER_VALUE,max_col=RECEIVER_HEADER + 3,fill=RECEIVER_FILL)
        }

        # ------------------------------------- Header Def -------------------------------------

        self.hr_header_range = {
            "min_row": HIERARCHY_ROW,
            "max_row": HIERARCHY_ROW,
            "min_col": self.init_col,
            "max_col": HR_MAX_COL
        }


